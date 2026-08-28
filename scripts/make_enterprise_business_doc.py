from __future__ import annotations

import sys
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT_DIR = ROOT / "deliverables"
OUT_PATH = OUT_DIR / "企业号产品业务手册_运营版_v1.0.docx"
TABLE_HELPERS = Path(
    "/Users/yuzijiang/.codex/plugins/cache/openai-primary-runtime/documents/26.819.11345/skills/documents/scripts"
)
sys.path.insert(0, str(TABLE_HELPERS))
from table_geometry import apply_table_geometry  # noqa: E402


# compact_reference_guide tokens; Arial Unicode MS is the named CJK-document
# override. Render QA exposes macOS system fonts through SAL_FONTPATH.
LATIN_FONT = "Arial Unicode MS"
CJK_FONT = "Arial Unicode MS"
NAVY = "17365D"
BLUE = "2B59C3"
DEEP_BLUE = "1F4D78"
INK = "222222"
MUTED = "667085"
LINE = "D7DFEA"
TABLE_HEAD = "E8EEF5"
LIGHT_BLUE = "F1F5FB"
LIGHT_GRAY = "F5F7FA"
LIGHT_GOLD = "FFF8E8"
GOLD = "9A6B00"
LIGHT_RED = "FFF1F0"
RED = "9B1C1C"
GREEN = "237B4B"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_border(cell, color: str = LINE, size: str = "6") -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.find(qn("w:tcBorders"))
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = qn(f"w:{edge}")
        node = borders.find(tag)
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), size)
        node.set(qn("w:color"), color)


def set_run_font(run, size: float | None = None, bold: bool | None = None,
                 color: str | None = None, italic: bool | None = None) -> None:
    run.font.name = LATIN_FONT
    run._element.get_or_add_rPr()
    fonts = run._element.rPr.rFonts
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        run._element.rPr.insert(0, fonts)
    fonts.set(qn("w:ascii"), LATIN_FONT)
    fonts.set(qn("w:hAnsi"), LATIN_FONT)
    fonts.set(qn("w:eastAsia"), CJK_FONT)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)


def configure_style(style, size: float, color: str, bold: bool,
                    before: float, after: float, line: float) -> None:
    style.font.name = LATIN_FONT
    style._element.get_or_add_rPr()
    fonts = style._element.rPr.rFonts
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        style._element.rPr.insert(0, fonts)
    fonts.set(qn("w:ascii"), LATIN_FONT)
    fonts.set(qn("w:hAnsi"), LATIN_FONT)
    fonts.set(qn("w:eastAsia"), CJK_FONT)
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.color.rgb = RGBColor.from_string(color)
    pf = style.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.line_spacing = line
    pf.keep_with_next = True


def configure_styles(doc: Document) -> None:
    normal = doc.styles["Normal"]
    configure_style(normal, 11, INK, False, 0, 6, 1.25)
    normal.paragraph_format.keep_with_next = False
    configure_style(doc.styles["Heading 1"], 16, BLUE, True, 18, 10, 1.1)
    configure_style(doc.styles["Heading 2"], 13, BLUE, True, 14, 7, 1.1)
    configure_style(doc.styles["Heading 3"], 12, DEEP_BLUE, True, 10, 5, 1.1)


def add_numbering_defs(doc: Document) -> tuple[int, int]:
    numbering = doc.part.numbering_part.element

    def next_id(tag: str, attr: str) -> int:
        vals = []
        for node in numbering.findall(qn(tag)):
            raw = node.get(qn(attr))
            if raw is not None and raw.isdigit():
                vals.append(int(raw))
        return max(vals, default=0) + 1

    def create(kind: str) -> int:
        abstract_id = next_id("w:abstractNum", "w:abstractNumId")
        abstract = OxmlElement("w:abstractNum")
        abstract.set(qn("w:abstractNumId"), str(abstract_id))
        multi = OxmlElement("w:multiLevelType")
        multi.set(qn("w:val"), "singleLevel")
        abstract.append(multi)
        lvl = OxmlElement("w:lvl")
        lvl.set(qn("w:ilvl"), "0")
        start = OxmlElement("w:start")
        start.set(qn("w:val"), "1")
        lvl.append(start)
        num_fmt = OxmlElement("w:numFmt")
        num_fmt.set(qn("w:val"), "bullet" if kind == "bullet" else "decimal")
        lvl.append(num_fmt)
        lvl_text = OxmlElement("w:lvlText")
        lvl_text.set(qn("w:val"), "•" if kind == "bullet" else "%1.")
        lvl.append(lvl_text)
        lvl_jc = OxmlElement("w:lvlJc")
        lvl_jc.set(qn("w:val"), "left")
        lvl.append(lvl_jc)
        p_pr = OxmlElement("w:pPr")
        tabs = OxmlElement("w:tabs")
        tab = OxmlElement("w:tab")
        tab.set(qn("w:val"), "num")
        tab.set(qn("w:pos"), "540")
        tabs.append(tab)
        p_pr.append(tabs)
        ind = OxmlElement("w:ind")
        ind.set(qn("w:left"), "540")
        ind.set(qn("w:hanging"), "270")
        p_pr.append(ind)
        spacing = OxmlElement("w:spacing")
        spacing.set(qn("w:after"), "80")
        spacing.set(qn("w:line"), "300")
        spacing.set(qn("w:lineRule"), "auto")
        p_pr.append(spacing)
        lvl.append(p_pr)
        r_pr = OxmlElement("w:rPr")
        r_fonts = OxmlElement("w:rFonts")
        r_fonts.set(qn("w:ascii"), LATIN_FONT)
        r_fonts.set(qn("w:hAnsi"), LATIN_FONT)
        r_fonts.set(qn("w:eastAsia"), CJK_FONT)
        r_pr.append(r_fonts)
        lvl.append(r_pr)
        abstract.append(lvl)
        numbering.append(abstract)

        num_id = next_id("w:num", "w:numId")
        num = OxmlElement("w:num")
        num.set(qn("w:numId"), str(num_id))
        abstract_num_id = OxmlElement("w:abstractNumId")
        abstract_num_id.set(qn("w:val"), str(abstract_id))
        num.append(abstract_num_id)
        numbering.append(num)
        return num_id

    return create("bullet"), create("decimal")


def apply_num(paragraph, num_id: int) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    num_pr = p_pr.find(qn("w:numPr"))
    if num_pr is None:
        num_pr = OxmlElement("w:numPr")
        p_pr.append(num_pr)
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    num_id_el = OxmlElement("w:numId")
    num_id_el.set(qn("w:val"), str(num_id))
    num_pr.append(ilvl)
    num_pr.append(num_id_el)


def add_bullet(doc: Document, text: str, bullet_id: int) -> None:
    p = doc.add_paragraph()
    apply_num(p, bullet_id)
    r = p.add_run(text)
    set_run_font(r, 11, color=INK)


def add_step(doc: Document, text: str, decimal_id: int) -> None:
    p = doc.add_paragraph()
    apply_num(p, decimal_id)
    r = p.add_run(text)
    set_run_font(r, 11, color=INK)


def restart_numbering(doc: Document, base_num_id: int) -> int:
    numbering = doc.part.numbering_part.element
    base = None
    for node in numbering.findall(qn("w:num")):
        if node.get(qn("w:numId")) == str(base_num_id):
            base = node
            break
    if base is None:
        return base_num_id
    abstract = base.find(qn("w:abstractNumId"))
    if abstract is None:
        return base_num_id
    ids = [int(node.get(qn("w:numId"))) for node in numbering.findall(qn("w:num"))
           if (node.get(qn("w:numId")) or "").isdigit()]
    new_id = max(ids, default=0) + 1
    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(new_id))
    abstract_ref = OxmlElement("w:abstractNumId")
    abstract_ref.set(qn("w:val"), abstract.get(qn("w:val")))
    num.append(abstract_ref)
    override = OxmlElement("w:lvlOverride")
    override.set(qn("w:ilvl"), "0")
    start = OxmlElement("w:startOverride")
    start.set(qn("w:val"), "1")
    override.append(start)
    num.append(override)
    numbering.append(num)
    return new_id


def add_steps(doc: Document, texts: list[str], base_decimal_id: int) -> None:
    num_id = restart_numbering(doc, base_decimal_id)
    for text in texts:
        add_step(doc, text, num_id)


def add_para(doc: Document, text: str, *, bold_lead: str | None = None,
             color: str = INK, after: float = 6, keep: bool = False) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.25
    p.paragraph_format.keep_with_next = keep
    if bold_lead and text.startswith(bold_lead):
        r1 = p.add_run(bold_lead)
        set_run_font(r1, 11, True, color)
        r2 = p.add_run(text[len(bold_lead):])
        set_run_font(r2, 11, False, color)
    else:
        r = p.add_run(text)
        set_run_font(r, 11, False, color)


def add_callout(doc: Document, title: str, text: str, *, fill: str = LIGHT_BLUE,
                title_color: str = BLUE) -> None:
    table = doc.add_table(rows=1, cols=1)
    apply_table_geometry(table, [9360], table_width_dxa=9360, indent_dxa=180,
                         cell_margins_dxa={"top": 150, "bottom": 150, "start": 180, "end": 180})
    cell = table.cell(0, 0)
    set_cell_shading(cell, fill)
    set_cell_border(cell, LINE, "8")
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run(title)
    set_run_font(r, 11, True, title_color)
    p2 = cell.add_paragraph()
    p2.paragraph_format.space_after = Pt(0)
    p2.paragraph_format.line_spacing = 1.2
    r2 = p2.add_run(text)
    set_run_font(r2, 10.5, False, INK)
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(2)
    spacer.paragraph_format.line_spacing = 0.5


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def add_table(doc: Document, headers: list[str], rows: list[list[str]], widths: list[int]) -> None:
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    hdr = table.rows[0]
    set_repeat_table_header(hdr)
    for i, text in enumerate(headers):
        cell = hdr.cells[i]
        set_cell_shading(cell, TABLE_HEAD)
        set_cell_border(cell)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(text)
        set_run_font(r, 10, True, NAVY)
    for row in rows:
        cells = table.add_row().cells
        for i, text in enumerate(row):
            cell = cells[i]
            set_cell_border(cell)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.15
            r = p.add_run(text)
            set_run_font(r, 9.5, False, INK)
    apply_table_geometry(table, widths, table_width_dxa=9360, indent_dxa=120,
                         cell_margins_dxa={"top": 100, "bottom": 100, "start": 120, "end": 120})
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def add_page_field(paragraph) -> None:
    run = paragraph.add_run()
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_char1, instr, fld_char2])
    set_run_font(run, 9, False, MUTED)


def configure_page(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)
    section.different_first_page_header_footer = True

    hp = section.header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.LEFT
    hp.paragraph_format.space_after = Pt(0)
    r = hp.add_run("企业号｜运营业务手册")
    set_run_font(r, 9, True, MUTED)
    fp = section.footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r2 = fp.add_run("内部业务参考  ·  ")
    set_run_font(r2, 9, False, MUTED)
    add_page_field(fp)


def add_cover(doc: Document) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(56)
    p.paragraph_format.space_after = Pt(8)
    r = p.add_run("OPERATIONS ENABLEMENT")
    set_run_font(r, 10, True, GOLD)

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(10)
    r = p.add_run("企业号产品业务手册")
    set_run_font(r, 30, True, NAVY)

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(28)
    r = p.add_run("面向运营的现有产品全景、业务流程、权限边界与问题处理口径")
    set_run_font(r, 13.5, False, MUTED)

    add_table(
        doc,
        ["文档属性", "内容"],
        [
            ["适用对象", "平台运营、商家运营、客服、培训与交付人员"],
            ["产品范围", "企业主端、普通员工端、小组长/组员协作、任务与收益闭环"],
            ["版本口径", "基于当前工作区主版本与 v61 已落地规则整理"],
            ["重要变更", "已移除企业主自定义用户昵称；用户昵称仅由用户本人维护"],
        ],
        [1700, 7660],
    )
    add_callout(
        doc,
        "一句话理解企业号",
        "企业号是围绕“企业组织协作—任务分发—内容生产与回填—订单收益—员工结算”搭建的多角色经营工作台。",
        fill=LIGHT_BLUE,
    )
    doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p.paragraph_format.space_before = Pt(36)
    r = p.add_run("产品业务文档  V1.0  ·  2026-08-24")
    set_run_font(r, 10, False, MUTED)
    doc.add_page_break()


def add_contents(doc: Document) -> None:
    doc.add_heading("阅读导航", level=1)
    entries = [
        ("01", "产品定位与业务闭环", "先建立企业号的整体认知"),
        ("02", "角色与权限边界", "理解企业主、普通员工、小组长、组员的差异"),
        ("03", "企业入驻与员工加入", "掌握从邀请到正式协作的完整流程"),
        ("04", "企业主经营工作台", "了解总览、运营中心、企业设置与记录管理"),
        ("05", "任务、关键词与内容回填", "理解核心生产链路及账号资产"),
        ("06", "收益、钱包与员工离职", "理解分佣、提现、离职回收与对账"),
        ("07", "运营日常与异常处理", "直接用于日常巡检、答疑和工单定位"),
        ("08", "口径、指标与范围边界", "统一对外表达，避免能力误读"),
    ]
    for num, title, desc in entries:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(9)
        r1 = p.add_run(f"{num}  ")
        set_run_font(r1, 12, True, BLUE)
        r2 = p.add_run(title)
        set_run_font(r2, 12, True, NAVY)
        r3 = p.add_run(f"\n      {desc}")
        set_run_font(r3, 10, False, MUTED)
    add_callout(
        doc,
        "运营先记住的三件事",
        "① 企业主统一管理员工身份、权限、项目与分成；② 配置变更通常只影响保存后的新订单，历史订单不追溯；③ 企业主不能替用户设置或修改昵称。",
        fill=LIGHT_GOLD,
        title_color=GOLD,
    )
    doc.add_page_break()


def build_document() -> Path:
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    doc = Document()
    configure_page(doc)
    configure_styles(doc)
    bullet_id, decimal_id = add_numbering_defs(doc)
    add_cover(doc)
    add_contents(doc)

    doc.add_heading("1. 产品定位与业务闭环", level=1)
    add_para(doc, "企业号服务于以项目推广、内容生产和订单结算为核心的企业协作场景。企业主在一个工作台中完成品牌配置、人员组织、项目权限、任务记录、收益查看和离职处置；员工在被授权的范围内参与任务、管理平台账号、提交关键词和作品回填，并查看个人收益或申请提现。")
    doc.add_heading("1.1 核心业务对象", level=2)
    add_table(
        doc,
        ["业务对象", "它是什么", "运营关注点"],
        [
            ["企业", "承载品牌、邀请码、项目授权、员工关系和经营数据的主体", "认证与品牌信息、可用项目、协议合规"],
            ["员工关系", "用户加入某企业后形成的协作关系，包含状态、角色、权限、分成和归属", "在职/离职、角色层级、收益与项目权限"],
            ["项目/任务", "企业被授权运营的业务项目，以及员工可参与的任务内容", "项目是否开放、任务规则、动态表单"],
            ["关键词", "员工参与任务后形成的可追踪生产资产，可继续提词、回填或转赠", "归属、状态、离职回收、转赠记录"],
            ["平台账号", "员工用于内容发布与回填的平台账号资料", "账号 ID/UID、主页链接、账号映射与可用性"],
            ["订单与收益", "回填结果经审核和订单拉取后形成的结算记录", "订单快照、分佣、结算状态、流水与提现"],
        ],
        [1500, 4280, 3580],
    )
    doc.add_heading("1.2 业务闭环", level=2)
    add_steps(doc, [
        "企业主完成企业品牌、邀请规则和项目预设配置，并生成二维码或邀请链接。",
        "用户登录、同意协议并提交加入申请；企业主审核通过后建立员工关系。",
        "企业主设置员工状态、真实角色、展示标签、项目权限、分成与收益查看权限。",
        "员工浏览任务并参与，形成关键词或组合提词记录；必要时维护发布平台账号。",
        "员工发布内容后进行单条或批量回填，平台完成审核、订单拉取和结算。",
        "收益按订单形成时的员工状态、身份、归属及分佣配置快照入账。",
        "员工完成实名认证、绑定收款账户并在授权范围内申请提现；企业主查看经营记录和整体收益。",
        "员工离职时，未提现余额自动回收至企业主，未完成关键词进入转赠池继续分配。",
    ], decimal_id)
    add_callout(doc, "配置生效原则", "员工身份、层级归属和分佣配置保存后立即生效，但只用于后续新订单；历史订单沿用形成时的快照，不做追溯重算。")

    doc.add_heading("2. 角色与权限边界", level=1)
    add_table(
        doc,
        ["角色", "核心职责", "典型可见能力", "主要限制"],
        [
            ["企业主", "经营管理与组织治理", "总览、员工/项目权限、任务与财务记录、邀请、品牌、离职处置", "不能替用户修改昵称；不应改写历史订单"],
            ["普通员工", "独立参与项目并提交成果", "任务、词库、账号、回填、个人收益、钱包", "仅访问企业和本人被授权的范围"],
            ["小组长", "带领组员并承担团队协作", "工作室、组员数据；经授权可看下属订单、个人收益和做单统计", "是否可看下属数据由企业主开关"],
            ["组员", "归属某位小组长参与任务", "员工侧任务、账号、回填、个人收益", "同一时间只能归属一名在职小组长"],
        ],
        [1250, 2250, 3550, 2310],
    )
    doc.add_heading("2.1 企业主管理能力", level=2)
    for text in [
        "管理员工在职/离职状态、收益数据查看权限和常规分成比例。",
        "设置真实身份为普通员工、小组长或组员，并维护小组长与组员的归属关系。",
        "配置员工可参与项目、字节任务台入口和提现权限；支持部分批量操作。",
        "查看企业总览、员工数据、推广/回填/收益/提现记录，并发放虚拟订单。",
        "处理员工离职后的余额回收和关键词重新分配。",
    ]:
        add_bullet(doc, text, bullet_id)

    doc.add_heading("2.2 昵称与角色标签：必须区分", level=2)
    add_callout(
        doc,
        "本版需求调整：移除企业主自定义用户昵称",
        "企业主端不提供“为员工设置昵称、备注昵称、企业内别名”能力，也不向员工配置接口提交昵称字段。列表、搜索、订单和记录中显示的昵称来自用户个人资料。",
        fill=LIGHT_RED,
        title_color=RED,
    )
    add_table(
        doc,
        ["字段/能力", "现有产品口径"],
        [
            ["用户昵称", "由用户本人在“账户中心—修改昵称”维护；企业主只能查看和按昵称搜索。"],
            ["企业主自定义用户昵称", "不在现有需求范围，产品不提供入口、字段、批量设置或企业内独立昵称。"],
            ["自定义角色标签", "继续保留。企业主可设置 1—5 个汉字，仅替换员工姓名旁的角色展示，不改变真实身份、权限、归属或分佣。"],
            ["协作识别", "运营定位人员时同时使用用户昵称、手机号、用户/协作 ID；昵称不应作为唯一身份凭证。"],
        ],
        [2500, 6860],
    )
    add_para(doc, "运营答复建议：用户反馈“企业主改不了员工昵称”时，应说明这是产品权限边界，并引导员工本人进入账户中心修改；若企业只想区分岗位，可使用自定义角色标签。")

    doc.add_heading("3. 企业入驻与员工加入", level=1)
    doc.add_heading("3.1 企业准备", level=2)
    add_steps(doc, [
        "配置员工端品牌名称、Logo 和简介，并确认企业可用协议；员工端采用白标展示。",
        "确认企业已获授权的项目范围，并在邀请设置中配置新员工默认的收益查看权限。",
        "设置项目预设：决定新加入员工默认开放哪些项目；后续也可同步给全部协作者。",
        "生成并分发二维码或邀请链接。",
    ], decimal_id)
    add_callout(doc, "预设生效范围", "邀请设置中的员工权限预设仅对配置变更后新加入的员工生效；项目预设可通过“同步给所有协作者”主动覆盖现有协作者的项目开关。", fill=LIGHT_GOLD, title_color=GOLD)

    doc.add_heading("3.2 用户加入与审核", level=2)
    add_steps(doc, [
        "用户通过二维码/链接进入邀请页，完成手机号验证码或密码登录。",
        "用户确认企业信息，阅读并同意相关协议后提交加入申请。",
        "企业主在“加入申请”查看申请人、手机号、申请时间等信息，并进行通过或驳回。",
        "驳回时填写原因，便于申请人补充后重新提交；已处理申请不可重复审核。",
        "通过后形成员工关系，企业主继续完成员工状态、项目、角色、分成等配置。",
    ], decimal_id)
    doc.add_heading("3.3 加入后的首轮检查", level=2)
    for text in [
        "员工是否为在职状态。",
        "员工是否看得到应参与的项目和任务。",
        "收益数据、提现、字节任务台等入口是否按运营方案开启。",
        "若启用进阶管理，小组长/组员关系和分成是否保存成功。",
        "员工本人昵称、手机号和协作 ID 是否能用于准确识别。",
    ]:
        add_bullet(doc, text, bullet_id)

    doc.add_heading("4. 企业主经营工作台", level=1)
    doc.add_heading("4.1 总览", level=2)
    add_para(doc, "总览用于快速判断企业经营规模和近期表现，包含企业主总收益、推广总收益、员工总收益、员工数量、推广任务量和回填量，并可按时间区间查看明细。员工维度可展示角色、收益、推广任务和回填数量。")
    add_callout(doc, "运营用途", "总览适合做健康度判断和异常发现，不替代订单/流水对账；金额问题要继续进入收益记录、订单记录或钱包流水核对。")

    doc.add_heading("4.2 运营中心", level=2)
    add_table(
        doc,
        ["模块", "企业主可以做什么", "运营重点"],
        [
            ["员工管理", "搜索/筛选员工，配置状态、权限、分成、角色标签、进阶角色与项目权限", "变更对新订单生效；离职需二次确认"],
            ["批量操作", "批量设置展示标签、分成、提现或字节任务台等支持的权限", "先核对选择范围，避免误操作"],
            ["虚拟订单", "选择员工、金额和关键词后生成业务订单", "确认员工与关键词有效，关注记录留痕"],
            ["关键词转赠", "处理离职关键词公池，将未完成关键词重新分配给在职员工", "仅更新当前归属，保留转赠记录"],
            ["记录中心", "查看推广任务、回填、收益和提现记录，支持关键词/昵称/项目等搜索", "按业务单号和人员信息定位异常"],
        ],
        [1500, 4700, 3160],
    )
    doc.add_heading("4.3 员工基础与进阶管理", level=2)
    add_para(doc, "基础管理包含员工状态、收益数据查看和常规分成。项目权限按项目逐一或批量开启/关闭。进阶管理开启后，普通员工可升级为小组长或组员，并配置层级归属、个人分成、团队业绩分成及下属数据查看权限。")
    for text in [
        "设置为组员时，必须选择一名在职小组长；组员只能归属一个小组长。",
        "设置为小组长时，可关联符合条件的在职普通员工，并可解除单个或多个组员关系。",
        "关闭进阶管理后恢复普通员工并解除层级关系。",
        "角色标签只是展示层；真实角色以系统保存的普通员工/小组长/组员为准。",
    ]:
        add_bullet(doc, text, bullet_id)

    doc.add_heading("4.4 企业设置与消息合规", level=2)
    add_para(doc, "“我”页聚合企业收益卡、快捷入口、消息通知和设置。企业设置支持员工端品牌名称、Logo、简介及合规协议查看；企业主还可获取员工邀请码、保存邀请配置和同步项目预设。平台公告、帮助中心、风险与处罚、税务与发票等内容用于运营解释和合规触达。")

    doc.add_heading("5. 任务、关键词与内容回填", level=1)
    doc.add_heading("5.1 员工任务链路", level=2)
    add_steps(doc, [
        "在员工首页/工作室查看被授权且当前可参与的任务。",
        "进入任务详情，阅读项目介绍、说明要求、流程、注意事项、教学和参与记录。",
        "点击参与后填写项目动态表单；不同项目可能要求书籍、关键词、账号或其他字段。",
        "提交后在“我的词库”查看关键词状态，并使用提词、组合提词或生成记录继续生产内容。",
        "发布内容后进入立即回填或批量回填，填写视频链接、平台账号、发布日期、播放量、附件等项目字段。",
        "提交后查看作品/回填记录和审核结果；通过后进入订单拉取与收益结算。",
    ], decimal_id)
    doc.add_heading("5.2 关键词生命周期", level=2)
    add_para(doc, "关键词是连接任务参与、内容生产、作品回填和收益追踪的关键资产。员工可在词库中查看状态、编辑被驳回记录、切换符合条件的项目、删除允许删除的记录或进行批量回填。组合提词可多选项目并形成组合回填。")
    add_callout(doc, "离职关键词规则", "员工离职后，其未完成关键词自动进入待转赠池。企业主选择新的在职员工后完成转赠；系统只更新当前归属并新增一条归属变更记录，不抹除原员工和操作轨迹。")

    doc.add_heading("5.3 发布账号资产", level=2)
    for text in [
        "员工可新增、编辑和删除平台账号，字段包括平台、账号 ID、账号名称、UID、DID、粉丝量和主页链接等。",
        "支持粘贴主页链接进行智能识别，也可手动添加；平台字段规则按项目和平台动态展示。",
        "CapCut 与 TikTok 场景可能要求账号一一对应，回填时需选择有效账号。",
        "账号昵称属于发布平台账号资料，不等同于企业主为员工设置用户昵称。",
    ]:
        add_bullet(doc, text, bullet_id)

    doc.add_heading("5.4 字节任务台", level=2)
    add_para(doc, "企业主可为员工开启字节任务台入口并查看员工绑定状态。员工进入后可获取机构专属链接，填写字节 UID、登录手机号并上传能显示机构名称和 UID 的截图。未获入口权限时，页面提示联系企业主开启。")

    doc.add_heading("6. 收益、钱包与员工离职", level=1)
    doc.add_heading("6.1 收益形成与展示", level=2)
    add_para(doc, "企业收益来自通过业务链路形成的订单。系统区分待结算与已结算，并提供订单日期、关键词、员工、项目、分佣比例和收益等明细。企业主可查看企业整体、推广、员工及个人维度的收益；员工是否可见收益页由企业主权限控制。")
    for text in [
        "普通员工订单按订单形成时保存的员工分成规则入账，企业主获得相应企业收益。",
        "进阶角色订单按保存时的小组长/组员归属、个人分成和团队业绩分成规则计算。",
        "身份、归属和分成调整只影响新订单；历史订单使用原快照。",
        "订单拉取时固化在职/离职状态，避免员工离职后仍新增待提现收益。",
    ]:
        add_bullet(doc, text, bullet_id)

    doc.add_heading("6.2 员工提现", level=2)
    add_steps(doc, [
        "员工完成实名认证；未认证、审核中或驳回时按页面提示处理。",
        "绑定个人收款账户/银行卡，并确保账户信息有效。",
        "企业主已为员工开启提现权限；员工设置提现密码。",
        "员工输入提现金额，系统计算预扣税额并提交申请。",
        "员工可查看审核中、已打款、已驳回等记录，以及实际到账、税额和时间信息。",
    ], decimal_id)

    doc.add_heading("6.3 离职收益回收", level=2)
    add_callout(
        doc,
        "当前正式规则",
        "员工有待提现余额也可以离职。企业主将状态从在职改为离职并二次确认后，系统一次性完成员工负数回收流水、企业主等额正数流水和离职回收累计金额更新。",
        fill=LIGHT_RED,
        title_color=RED,
    )
    for text in [
        "企业主在员工管理中可直接看到当前待提现金额；切换离职后保存按钮变为“确认离职并保存”。",
        "确认弹层明确展示员工、余额和三步处理结果；取消不提交。",
        "回收后员工待提现金额归零，企业主“我”页收益卡增加“离职员工回收金额”。",
        "员工负流水与企业主正流水金额一致，并可通过同一回收配对号进行对账；重复离职请求不得重复回收。",
        "离职后的新订单不得再进入员工待提现，按服务端既定离职订单归属规则计入企业主。",
    ]:
        add_bullet(doc, text, bullet_id)
    doc.add_heading("6.4 离职对账检查", level=2)
    for text in [
        "员工状态已变为离职，待提现金额为 0。",
        "员工出现一笔等额负数离职回收流水。",
        "企业主出现一笔等额正数离职员工收益回收流水。",
        "两笔流水回收配对号一致，离职回收累计金额同步增加。",
        "后续订单未再增加该员工待提现；未完成关键词已进入转赠池。",
    ]:
        add_bullet(doc, text, bullet_id)

    doc.add_heading("7. 运营日常与异常处理", level=1)
    doc.add_heading("7.1 日常巡检清单", level=2)
    add_table(
        doc,
        ["频率", "建议动作", "关注结果"],
        [
            ["每日", "查看加入申请、消息公告、回填/收益/提现异常记录", "申请及时处理，失败或驳回有明确原因"],
            ["每日", "关注员工离职、余额回收和关键词待转赠", "余额只回收一次，关键词及时重新分配"],
            ["每周", "查看企业总览、员工活跃与任务/回填趋势", "定位低参与、低回填或收益异常员工"],
            ["每周", "抽查角色、项目、收益查看、提现和字节入口权限", "权限与运营方案一致，无批量误配"],
            ["配置后", "验证新员工预设、项目同步及新订单快照", "生效范围与时间符合预期"],
        ],
        [1250, 4930, 3180],
    )

    doc.add_heading("7.2 常见问题定位", level=2)
    cases = [
        ("员工看不到任务", "先确认员工在职；再检查企业是否获项目授权、员工项目权限、邀请预设/同步结果和任务当前状态。"),
        ("员工看不到收益", "检查员工的“收益数据”权限；确认菜单由服务端按角色和权限返回。"),
        ("员工不能提现", "依次核对实名认证、收款账户、企业主提现权限、提现密码、可提现余额及税费计算。"),
        ("企业主想改员工昵称", "明确告知不支持。由员工本人在账户中心修改昵称；岗位区分使用自定义角色标签。"),
        ("离职后余额没有归零", "检查离职是否真正确认成功、员工/企业主是否出现成对流水及累计回收金额；按员工 ID、回收配对号升级对账。"),
        ("离职后员工仍有新收益", "核对订单日期、订单拉取时间和在职状态快照；历史订单不追溯，新订单不应进入离职员工待提现。"),
        ("关键词无法转赠", "确认关键词位于离职待转赠池、接收人是在职且符合候选条件，并核对转赠记录。"),
        ("字节任务台未开放", "由企业主开启员工入口；员工再提交 UID、手机号和两张要求截图。"),
        ("回填识别失败", "检查链接平台和格式、账号对应关系、项目字段是否完整；无法识别时改为手动回填。"),
    ]
    for title, answer in cases:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.keep_with_next = True
        r = p.add_run(title)
        set_run_font(r, 11, True, NAVY)
        p2 = doc.add_paragraph()
        p2.paragraph_format.left_indent = Inches(0.18)
        p2.paragraph_format.space_after = Pt(7)
        p2.paragraph_format.line_spacing = 1.2
        r2 = p2.add_run(answer)
        set_run_font(r2, 10.5, False, INK)

    doc.add_heading("7.3 工单最小信息集", level=2)
    for text in [
        "企业 ID/名称、用户 ID、员工关系 ID及手机号后四位。",
        "问题发生时间、入口路径、角色和当前员工状态。",
        "项目 ID、关键词 ID、订单号、提现号或回收配对号等业务标识。",
        "页面截图、提示文案、是否可稳定复现。",
        "涉及金额时同时提供员工流水与企业主流水，不仅提供汇总数。",
    ]:
        add_bullet(doc, text, bullet_id)

    doc.add_heading("8. 口径、指标与范围边界", level=1)
    doc.add_heading("8.1 核心指标口径", level=2)
    add_table(
        doc,
        ["指标", "业务含义", "使用提醒"],
        [
            ["企业主总收益", "企业主维度累计收益", "金额异常需下钻订单和钱包流水"],
            ["推广总收益", "企业推广业务产生的整体收益", "与员工收益并非同义字段"],
            ["员工总收益", "企业员工维度汇总收益", "受订单状态、分佣和员工状态影响"],
            ["待结算", "已形成但尚未完成结算的收益", "不是可直接提现金额"],
            ["已结算", "已完成结算的收益", "仍需结合钱包可提现口径"],
            ["待提现金额", "员工当前可纳入提现或离职回收的余额", "离职确认后应清零"],
            ["离职员工回收金额", "企业主累计收到的离职员工待提现余额回收", "应等于有效企业主正向回收流水之和"],
        ],
        [2100, 4300, 2960],
    )

    doc.add_heading("8.2 对外统一表述", level=2)
    for text in [
        "企业号支持企业统一管理员工身份、项目权限、分成和任务协作。",
        "员工数据和菜单会随角色、在职状态及权限配置变化。",
        "用户昵称属于个人资料，由用户本人维护；企业主不能自定义或覆盖员工昵称。",
        "企业主可自定义员工姓名旁的角色标签，但标签不改变真实权限和分佣。",
        "员工离职不会因存在待提现余额而被阻断，余额会自动成对回收并可对账。",
        "历史订单按原快照保留，角色或分佣变更只影响后续新订单。",
    ]:
        add_bullet(doc, text, bullet_id)

    doc.add_heading("8.3 当前范围内与范围外", level=2)
    add_table(
        doc,
        ["当前产品范围内", "本版明确不包含"],
        [
            ["企业品牌白标、邀请与申请审核", "企业主为员工设置企业内昵称/备注昵称"],
            ["员工状态、角色层级、展示标签和项目权限", "因配置变更自动重算历史订单"],
            ["任务、关键词、账号、回填、订单和收益闭环", "运营在无业务标识时仅凭昵称完成资金对账"],
            ["实名认证、收款账户、税费与提现记录", "绕过实名、权限或合规校验的人工提现"],
            ["离职余额回收和关键词转赠留痕", "删除离职、转赠或资金流水历史"],
        ],
        [4680, 4680],
    )
    add_callout(
        doc,
        "文档维护建议",
        "后续如新增角色、菜单、资金口径或审核状态，应同时更新角色矩阵、业务闭环、异常处理和指标口径；若再次讨论昵称能力，必须明确区分“用户本人昵称”与“企业角色标签”。",
        fill=LIGHT_GRAY,
        title_color=NAVY,
    )

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(20)
    r = p.add_run("— 完 —")
    set_run_font(r, 10, False, MUTED)

    # Core properties and document hygiene.
    props = doc.core_properties
    props.title = "企业号产品业务手册（运营版）"
    props.subject = "企业号现有产品业务、角色权限、核心流程与运营口径"
    props.author = "产品团队"
    props.keywords = "企业号, 运营手册, 企业主, 员工, 任务, 收益, 离职回收"
    props.comments = "已移除企业主自定义用户昵称能力口径"

    doc.save(OUT_PATH)
    return OUT_PATH


if __name__ == "__main__":
    path = build_document()
    print(path)
